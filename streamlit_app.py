from datetime import datetime
import os
import openpyxl
import pandas as pd
import streamlit as st

st.set_page_config(
    page_title="Campeonato de Marcas - Ganadores & Pulso", layout="wide"
)

# --- ENCABEZADO CON LOGOTIPO ---
col_logo1, col_logo2, col_logo3 = st.columns([1, 2, 1])
with col_logo2:
  # Asegúrate de subir tu imagen cuadrada con este nombre exacto a tu repositorio de GitHub
  if os.path.exists("logo.jpg"):
    st.image("logo.jpg", use_container_width=True)
  elif os.path.exists("Logo.jpg"):
    st.image("Logo.jpg", use_container_width=True)
  else:
    st.title("🏇 Campeonato de Marcas - Ganadores & Pulso")

st.markdown(
    "<h3 style='text-align: center; color: #4A90E2;'>Proyectando el"
    " Hipismo</h3>",
    unsafe_allow_html=True,
)
st.markdown(
    "<p style='text-align: center; font-style: italic; color: #888;'>Director:"
    " Nelson Osorio</p>",
    unsafe_allow_html=True,
)
st.markdown("---")

excel_file = "CONTROL CAMPEONATO DE MARCAS.xlsx"

# Inicializar variables globales de configuración en la sesión si no existen
if "semana_activa" not in st.session_state:
  st.session_state["semana_activa"] = "SEM 1"
if "carreras_activas" not in st.session_state:
  st.session_state["carreras_activas"] = 14
if "participante_logueado" not in st.session_state:
  st.session_state["participante_logueado"] = None

if os.path.exists(excel_file):
  try:
    # Cargamos el libro con openpyxl para leer coordenadas exactas
    wb_lee = openpyxl.load_workbook(excel_file, data_only=True)
    if "CONTROL" not in wb_lee.sheetnames:
      st.error("No se encontró la hoja 'CONTROL' en el archivo de Excel.")
      st.stop()

    ws_lee = wb_lee["CONTROL"]

    # Leemos los encabezados de la Fila 3, limitados hasta la columna R (18)
    max_col_permitida = 18
    headers = [
        ws_lee.cell(row=3, column=col).value
        for col in range(1, max_col_permitida + 1)
    ]
    headers_limpios = [
        str(h).strip() if h is not None else f"Col_{i+1}"
        for i, h in enumerate(headers)
    ]

    # Leemos los datos desde la Fila 4 hasta la Fila 53
    max_fila_permitida = min(53, ws_lee.max_row)
    datos_filas = []
    lista_participantes_validos = {}

    for r in range(4, max_fila_permitida + 1):
      val_item = ws_lee.cell(row=r, column=1).value
      val_codigo = ws_lee.cell(row=r, column=2).value
      val_nombre = ws_lee.cell(row=r, column=3).value

      if (
          val_codigo is None
          or str(val_codigo).strip() == ""
          or str(val_codigo).lower() in ["none", "nan"]
      ):
        continue
      if (
          val_nombre is None
          or str(val_nombre).strip() == ""
          or str(val_nombre).lower() in ["none", "nan"]
      ):
        continue

      codigo_str = str(val_codigo).strip()
      nombre_str = str(val_nombre).strip()

      lista_participantes_validos[nombre_str] = codigo_str

      fila_valores = [
          ws_lee.cell(row=r, column=c).value
          for c in range(1, max_col_permitida + 1)
      ]
      fila_valores_limpios = ["" if v is None else v for v in fila_valores]
      datos_filas.append(fila_valores_limpios)

    df_final = pd.DataFrame(datos_filas, columns=headers_limpios)

    # Menú Lateral Actualizado
    st.sidebar.header("🔐 Menú de Navegación")
    modo = st.sidebar.radio(
        "Seleccione una opción:",
        [
            "Acceso de Participantes",
            "Listado de Participantes",
            "Ranking General",
            "Panel de Director",
        ],
    )

    if modo == "Listado de Participantes":
      st.subheader("📋 Listado Oficial de Participantes")
      st.dataframe(df_final, use_container_width=True, hide_index=True)

    elif modo == "Ranking General":
      st.subheader("🏆 Ranking General del Campeonato")
      hoja_ranking = "20-9-26"

      if hoja_ranking in wb_lee.sheetnames:
        ws_rank = wb_lee[hoja_ranking]
        # Extraer desde celda B2 hasta Q151 (Columna B es 2, Columna Q es 17)
        filas_ranking = []
        for r in range(2, 152):
          fila_vals = [ws_rank.cell(row=r, column=c).value for c in range(2, 18)]
          if all(v is None or str(v).strip() == "" for v in fila_vals):
            continue
          filas_ranking.append(
              ["" if v is None else v for v in fila_vals]
          )

        if filas_ranking:
          raw_headers = [
              str(h).strip() if h is not None and str(h).strip() != "" else f"Col_{i+1}"
              for i, h in enumerate(filas_ranking[0])
          ]
          seen = {}
          unique_headers = []
          for h in raw_headers:
            if h in seen:
              seen[h] += 1
              unique_headers.append(f"{h}_{seen[h]}")
            else:
              seen[h] = 0
              unique_headers.append(h)

          datos_rank = filas_ranking[1:]
          df_rank = pd.DataFrame(datos_rank, columns=unique_headers)
          st.dataframe(df_rank, use_container_width=True, hide_index=True)
        else:
          st.warning(
              f"La hoja '{hoja_ranking}' no contiene datos en el rango"
              " especificado."
          )
      else:
        st.warning(
            f"⚠️ No se encontró la hoja '{hoja_ranking}' en el archivo de Excel."
        )

    elif modo == "Acceso de Participantes":
      st.subheader("✍️ Módulo de Ingreso de Marcas")

      if not st.session_state["participante_logueado"]:
        st.markdown(
            "Por favor ingrese su **Nombre** y su **Clave de Acceso** (su"
            " código asignado, ej: `COD 02`)."
        )

        with st.form("form_login"):
          nombre_ingresado = st.selectbox(
              "Seleccione su Nombre:",
              options=list(lista_participantes_validos.keys()),
          )
          clave_ingresada = st.text_input(
              "Clave de Seguridad (Código):", type="password"
          )
          btn_login = st.form_submit_button("🔑 Ingresar al Sistema")

          if btn_login:
            codigo_correcto = lista_participantes_validos.get(nombre_ingresado)
            if clave_ingresada.strip().upper() == codigo_correcto.upper():
              st.session_state["participante_logueado"] = {
                  "nombre": nombre_ingresado,
                  "codigo": codigo_correcto,
              }
              st.success(
                  f"¡Bienvenido, {nombre_ingresado} ({codigo_correcto})!"
              )
              st.rerun()
            else:
              st.error(
                  "❌ Clave incorrecta. Recuerde que su clave es su código"
                  " asignado (ej. COD 01)."
              )
      else:
        p_info = st.session_state["participante_logueado"]
        num_carr = st.session_state["carreras_activas"]
        limite_no_validas = num_carr - 6

        st.info(
            f"👤 Participante Activo: **{p_info['nombre']}** (`{p_info['codigo']}`)"
            f" | Jornada Activa: **{st.session_state['semana_activa']}** |"
            f" Carreras Habilitadas: **{num_carr}**"
        )

        if st.button("🚪 Cerrar Sesión / Cambiar de Participante"):
          st.session_state["participante_logueado"] = None
          st.rerun()

        st.markdown("---")
        st.markdown("### 📌 Reglamento Oficial de Marcas:")
        st.markdown(
            f"- **Carreras No Válidas (1 a {limite_no_validas}):** 3 marcas"
            " por carrera y exactamente **1 Fija (F)** obligatoria."
        )
        st.markdown(
            f"- **Válidas del 5y6 ({limite_no_validas + 1} a {num_carr}):** 3"
            " marcas por carrera y exactamente **1 Fija (F)** obligatoria."
        )
        st.markdown(
            "- **Súper Fijo (SF):** Solo se marca **1 ejemplar** en toda la"
            " reunión (las casillas 2 y 3 deben ir vacías en esa carrera)."
        )
        st.markdown("---")

        with st.form("form_marcas"):
          marcas_por_carrera = {}

          for i in range(1, num_carr + 1):
            st.markdown(f"#### 🏁 Carrera {i}")

            col_m1, col_m2, col_m3, col_m4 = st.columns([2, 2, 2, 2])

            with col_m1:
              m1 = st.text_input(f"1ra Marca C-{i}", key=f"c{i}_m1")
            with col_m2:
              m2 = st.text_input(f"2da Marca C-{i}", key=f"c{i}_m2")
            with col_m3:
              m3 = st.text_input(f"3ra Marca C-{i}", key=f"c{i}_m3")
            with col_m4:
              tipo_jugada = st.selectbox(
                  f"Tipo C-{i}",
                  ["Normal", "Fijo (F)", "Súper Fijo (SF)"],
                  key=f"c{i}_tipo",
              )

            marcas_por_carrera[i] = {
                "m1": m1,
                "m2": m2,
                "m3": m3,
                "tipo": tipo_jugada,
            }
            st.markdown("---")

          boton_guardar = st.form_submit_button("🚀 Guardar y Enviar Marcas")

          if boton_guardar:
            sf_count = sum(
                1
                for c in marcas_por_carrera.values()
                if c["tipo"] == "Súper Fijo (SF)"
            )
            error_val = False

            if sf_count > 1:
              st.error(
                  "⚠️ Solo puedes seleccionar un (1) Súper Fijo (SF) por reunión."
              )
              error_val = True
            else:
              for num, datos in marcas_por_carrera.items():
                if datos["tipo"] == "Súper Fijo (SF)" and (
                    datos["m2"].strip() != "" or datos["m3"].strip() != ""
                ):
                  st.error(
                      f"⚠️ En la Carrera {num} seleccionaste Súper Fijo (SF), la"
                      " 2da y 3ra marca deben ir vacías."
                  )
                  error_val = True
                  break

            if not error_val:
              fijas_no_val = sum(
                  1
                  for i in range(1, limite_no_validas + 1)
                  if marcas_por_carrera[i]["tipo"] == "Fijo (F)"
              )
              fijas_val = sum(
                  1
                  for i in range(limite_no_validas + 1, num_carr + 1)
                  if marcas_por_carrera[i]["tipo"] == "Fijo (F)"
              )

              if fijas_no_val != 1:
                st.error(
                    f"⚠️ Debes seleccionar exactamente una (1) Fija (F) en el"
                    f" bloque de carreras no válidas (Carreras 1 a"
                    f" {limite_no_validas}). Tienes {fijas_no_val} seleccionada(s)."
                )
                error_val = True
              elif fijas_val != 1:
                st.error(
                    f"⚠️ Debes seleccionar exactamente una (1) Fija (F) en el"
                    f" bloque de válidas del 5y6 (Carreras {limite_no_validas + 1}"
                    f" a {num_carr}). Tienes {fijas_val} seleccionada(s)."
                )
                error_val = True

            if not error_val:
              try:
                wb = openpyxl.load_workbook(excel_file)
                semana_actual = st.session_state["semana_activa"]

                if semana_actual in wb.sheetnames:
                  ws = wb[semana_actual]
                else:
                  ws = wb.create_sheet(title=semana_actual)
                  cabecera = ["NOMBRE PARTICIPANTE"]
                  for c in range(1, 17):
                    cabecera.extend([
                        f"C{c}_1ra",
                        f"C{c}_2da",
                        f"C{c}_3ra",
                        f"C{c}_Tipo",
                    ])
                  ws.append(cabecera)

                fila_encontrada = None
                for row_idx in range(4, ws.max_row + 1):
                  val_celda = ws.cell(row=row_idx, column=3).value
                  if (
                      val_celda
                      and str(val_celda).strip().upper()
                      == p_info["nombre"].strip().upper()
                  ):
                    fila_encontrada = row_idx
                    break

                datos_fila = [p_info["nombre"]]
                for c in range(1, 17):
                  if c <= num_carr:
                    m_data = marcas_por_carrera[c]
                    tipo_corto = "N"
                    if m_data["tipo"] == "Fijo (F)":
                      tipo_corto = "F"
                    elif m_data["tipo"] == "Súper Fijo (SF)":
                      tipo_corto = "SF"

                    datos_fila.extend([
                        m_data["m1"],
                        m_data["m2"],
                        m_data["m3"],
                        tipo_corto,
                    ])
                  else:
                    datos_fila.extend(["", "", "", ""])

                if fila_encontrada:
                  for col_idx, val in enumerate(datos_fila, start=1):
                    ws.cell(row=fila_encontrada, column=col_idx, value=val)
                else:
                  ws.append(datos_fila)

                wb.save(excel_file)
                st.success(
                    f"¡Tus marcas han sido guardadas con éxito en la hoja"
                    f" '{semana_actual}'!"
                )
              except Exception as err:
                st.error(f"Error al escribir en el Excel: {err}")

    elif modo == "Panel de Director":
      st.subheader("🛠️ Panel Privado del Director (Nelson Osorio)")
      st.markdown(
          "Configura la jornada activa y descarga el archivo actualizado con"
          " las marcas recibidas."
      )

      clave_director = st.text_input(
          "Contraseña de Administrador:", type="password"
      )
      if clave_director == "ganadores2026":
        st.success("✅ Acceso de Director Autorizado.")
        st.markdown("---")
        st.markdown("#### ⚙️ Configuración Global de la Jornada")

        nueva_semana = st.selectbox(
            "Seleccione la Semana Activa para los participantes:",
            ["SEM 1", "SEM 2", "SEM 3", "SEM 4", "SEM 5", "SEM 6"],
            index=["SEM 1", "SEM 2", "SEM 3", "SEM 4", "SEM 5", "SEM 6"].index(
                st.session_state["semana_activa"]
            ),
        )

        nuevas_carreras = st.slider(
            "Cantidad de carreras programadas para esta semana:",
            min_value=6,
            max_value=16,
            value=st.session_state["carreras_activas"],
        )

        if st.button("💾 Guardar Configuración de la Jornada"):
          st.session_state["semana_activa"] = nueva_semana
          st.session_state["carreras_activas"] = nuevas_carreras
          st.success(
              f"¡Configuración actualizada! Semana activa: **{nueva_semana}** |"
              f" Carreras: **{nuevas_carreras}**"
          )

        st.markdown("---")
        st.markdown("#### 📥 Descarga de Base de Datos")
        with open(excel_file, "rb") as f:
          st.download_button(
              label="📥 Descargar Excel Actualizado con Marcas",
              data=f,
              file_name="CONTROL CAMPEONATO DE MARCAS.xlsx",
              mime=(
                  "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
              ),
          )
      elif clave_director:
        st.error("❌ Contraseña incorrecta.")

  except Exception as e:
    st.error(f"Error al procesar el archivo de Excel: {e}")
else:
  st.error("No se encuentra el archivo de Excel en la carpeta.")
